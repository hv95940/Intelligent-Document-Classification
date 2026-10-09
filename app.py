import streamlit as st
import joblib
from pathlib import Path
from pypdf import PdfReader
from docx import Document

st.set_page_config(
    page_title="Intelligent Document Classification",
    page_icon="📄",
    layout="centered"
)

st.title("📄 Intelligent Document Classification System")
st.write("Upload a PDF, DOCX, or TXT document to classify it.")

# Load model
MODEL_PATH = Path(__file__).parent / "document_classifier.pkl"

try:
    model = joblib.load(MODEL_PATH)
except Exception as error:
    st.error(f"Unable to load model: {error}")
    st.stop()


# Extract document text
def extract_text(file):
    filename = file.name.lower()

    try:
        if filename.endswith(".pdf"):
            reader = PdfReader(file)
            return "\n".join(
                page.extract_text() or ""
                for page in reader.pages
            ).strip()

        elif filename.endswith(".docx"):
            document = Document(file)
            paragraphs = [
                paragraph.text
                for paragraph in document.paragraphs
            ]

            for table in document.tables:
                for row in table.rows:
                    paragraphs.append(
                        " ".join(cell.text for cell in row.cells)
                    )

            return "\n".join(paragraphs).strip()

        elif filename.endswith(".txt"):
            return file.getvalue().decode(
                "utf-8-sig", errors="replace"
            ).strip()

        return ""

    except Exception as error:
        st.error(f"Text extraction failed: {error}")
        return ""


# Upload file
uploaded_file = st.file_uploader(
    "Upload your document",
    type=["pdf", "docx", "txt"]
)

if uploaded_file is not None:
    st.success(f"Uploaded: {uploaded_file.name}")

    text = extract_text(uploaded_file)

    if text:
        st.subheader("Extracted Text")
        st.text_area(
            "Document content",
            value=text,
            height=200
        )

        if st.button("Predict Category"):
            try:
                prediction = model.predict([text])[0]
                st.success(f"Predicted Category: {prediction}")

                if hasattr(model, "predict_proba"):
                    probabilities = model.predict_proba([text])[0]
                    confidence = max(probabilities) * 100
                    st.info(f"Model confidence: {confidence:.2f}%")

            except Exception as error:
                st.error(f"Prediction failed: {error}")

    else:
        st.warning(
            "No readable text found in this file. "
            "Scanned PDFs may require OCR."
        )
else:
    st.info("Please upload a document to begin.")

